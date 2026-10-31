import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = ['<w:tcBorders ' + nsdecls("w") + '>']
    for side, border in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if border:
            borders.append(f'<w:{side} w:val="{border.get("val", "single")}" w:sz="{border.get("sz", "4")}" w:space="0" w:color="{border.get("color", "auto")}"/>')
        else:
            borders.append(f'<w:{side} w:val="none"/>')
    borders.append('</w:tcBorders>')
    tcPr.append(parse_xml(''.join(borders)))

def create_document():
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    
    # Configure Normal style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(9)
    font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return doc

def build_poster_docx_en(output_path):
    doc = create_document()
    
    # 1. Header Banner Table
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.8)
    header_table.columns[1].width = Inches(2.57)
    
    c0 = header_table.cell(0, 0)
    c1 = header_table.cell(0, 1)
    for c in (c0, c1):
        set_cell_background(c, "091A36")
        set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(2)
    run_badge = p0.add_run("ASSISTIVE AI & COMPUTER VISION RESEARCH PROJECT\n")
    run_badge.font.size = Pt(7)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    
    run_title = p0.add_run("APEC SMART GLASSES FOR THE BLIND\n")
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    run_sub = p0.add_run("Real-Time Multimodal Vision AI, Local Vietnamese ASR & Tactile Accessibility")
    run_sub.font.size = Pt(8.5)
    run_sub.font.color.rgb = RGBColor(0xBA, 0xE6, 0xFD)
    
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    
    meta_text = [
        ("Repository: ", "thanh1912-ut/APEC_mat_kinh..."),
        ("Target SHA: ", "d9ba1e0 (main branch)"),
        ("Current Release: ", "Android Pilot v0.4.0 (Code 8)"),
        ("Speech Engine: ", "Zipformer RNN-T (sherpa-onnx)"),
        ("Vision AI: ", "Google Gemini 3.8 Flash"),
    ]
    for lbl, val in meta_text:
        r_lbl = p1.add_run(lbl)
        r_lbl.font.size = Pt(6.8)
        r_lbl.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        r_val = p1.add_run(val + "\n")
        r_val.font.size = Pt(6.8)
        r_val.font.bold = True
        r_val.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # 2. Key Metrics Ribbon (4 Columns)
    metrics_table = doc.add_table(rows=1, cols=4)
    metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    metrics_table.autofit = False
    
    col_w = Inches(7.37 / 4.0)
    metrics = [
        ("< 180 ms", "Local STT Latency", "0284C7"),
        ("100% Tactile", "Hardware Volume Keys", "10B981"),
        ("Zero Audio Leak", "Speech Kept On-Premises", "8B5CF6"),
        ("651 Test Cases", "322 Python + 329 Jest", "F59E0B"),
    ]
    for i, (val, lbl, color) in enumerate(metrics):
        cell = metrics_table.cell(0, i)
        cell.width = col_w
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        set_cell_borders(cell, top={"val": "single", "sz": "24", "color": color},
                               bottom={"val": "single", "sz": "4", "color": "E2E8F0"},
                               left={"val": "single", "sz": "4", "color": "E2E8F0"},
                               right={"val": "single", "sz": "4", "color": "E2E8F0"})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        r_v = p.add_run(val + "\n")
        r_v.font.size = Pt(10)
        r_v.font.bold = True
        r_v.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_l = p.add_run(lbl)
        r_l.font.size = Pt(6.5)
        r_l.font.bold = True
        r_l.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3. Two-Column Content Layout
    layout_table = doc.add_table(rows=1, cols=2)
    layout_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    layout_table.autofit = False
    layout_table.columns[0].width = Inches(3.62)
    layout_table.columns[1].width = Inches(3.62)
    
    col_left = layout_table.cell(0, 0)
    col_right = layout_table.cell(0, 1)
    for c in (col_left, col_right):
        set_cell_margins(c, top=0, bottom=0, left=40, right=40)
    
    # Left Column: Card 1 - Mission & Goals
    p_c1_title = col_left.paragraphs[0]
    p_c1_title.paragraph_format.space_before = Pt(0)
    p_c1_title.paragraph_format.space_after = Pt(2)
    r1 = p_c1_title.add_run("1. PROJECT MISSION, GOALS & SCOPE")
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    p_c1_desc = col_left.add_paragraph()
    p_c1_desc.paragraph_format.space_after = Pt(3)
    p_c1_desc.paragraph_format.line_spacing = 1.15
    r_desc = p_c1_desc.add_run(
        "APEC Smart Glasses is an assistive engineering project dedicated to providing blind and visually impaired users with independent spatial awareness, instant object recognition, and interactive guidance for domestic appliances."
    )
    r_desc.font.size = Pt(7.4)
    
    bullets_c1 = [
        ("True Non-Visual Accessibility: ", "Eliminates touchscreen barriers by utilizing physical smartphone volume keys (Volume Up = Listen/Capture/Send; Volume Down = Cancel/Stop) paired with audio earcons."),
        ("Absolute Speech Privacy: ", "Streaming speech recognition executes entirely on-premises using lightweight Zipformer RNN-T models (~51MB ONNX). Raw user speech is never leaked to cloud APIs."),
        ("Appliance Guidance Engine (v0.4.0): ", "Provides verified step-by-step assistance to locate and trigger physical controls on air conditioners, TV remotes, and household appliances."),
        ("Spatial Vision Roadmap: ", "Pre-engineered abstraction layer designed to integrate STMicroelectronics VL53L9CX Direct ToF LiDAR sensors for 3D millimeter-level hazard detection.")
    ]
    for b_title, b_text in bullets_c1:
        p_b = col_left.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(2.5)
        p_b.paragraph_format.line_spacing = 1.15
        p_b.paragraph_format.left_indent = Inches(0.12)
        r_bt = p_b.add_run("■ " + b_title)
        r_bt.font.size = Pt(7.2)
        r_bt.font.bold = True
        r_bt.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_bx = p_b.add_run(b_text)
        r_bx.font.size = Pt(7.0)
        r_bx.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # Left Column: Card 2 - Architecture & Pipeline
    p_c2_title = col_left.add_paragraph()
    p_c2_title.paragraph_format.space_before = Pt(5)
    p_c2_title.paragraph_format.space_after = Pt(2)
    r2 = p_c2_title.add_run("2. SYSTEM ARCHITECTURE & DATA FLOW")
    r2.font.size = Pt(8.5)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    steps = [
        ("A. Mobile Client", "React Native/Expo app captures rear camera JPEG (max 1280px) and voice query (M4A/AAC)."),
        ("B. FastAPI Gateway", "Authenticates SHA-256 Bearer tokens, enforces 4MiB limits, queues into WorkGate."),
        ("C. Local ASR Worker", "PyAV resamples audio to mono 16kHz; sherpa-onnx transcribes Vietnamese text in RAM."),
        ("D. Vision AI", "Image and transcript queried to Gemini 3.8 Flash / 3.1 Flash Lite with store=false."),
        ("E. Audio Delivery", "Structured JSON response spoken via Android TTS and earcons without TalkBack clash.")
    ]
    for s_step, s_desc in steps:
        p_s = col_left.add_paragraph()
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = Pt(2)
        p_s.paragraph_format.line_spacing = 1.15
        p_s.paragraph_format.left_indent = Inches(0.1)
        r_st = p_s.add_run(f"[{s_step}]: ")
        r_st.font.size = Pt(7.0)
        r_st.font.bold = True
        r_st.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        r_sx = p_s.add_run(s_desc)
        r_sx.font.size = Pt(6.8)
    
    # Left Column: Spec Table
    t_spec = col_left.add_table(rows=6, cols=2)
    t_spec.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_spec.autofit = False
    t_spec.columns[0].width = Inches(1.3)
    t_spec.columns[1].width = Inches(2.2)
    specs = [
        ("Component", "Technology & Specification"),
        ("Mobile Client", "React Native, Expo SDK 57, TypeScript"),
        ("Backend Gateway", "Python 3.11, FastAPI, PyAV, WorkGate"),
        ("Vietnamese STT", "sherpa-onnx (Zipformer RNN-T 51MB)"),
        ("Multimodal Vision", "Google Gemini 3.8 Flash (store=false)"),
        ("Guidance State", "PostgreSQL + Redis (CAS Locks)")
    ]
    for r_idx, (c1_txt, c2_txt) in enumerate(specs):
        cell1 = t_spec.cell(r_idx, 0)
        cell2 = t_spec.cell(r_idx, 1)
        set_cell_margins(cell1, top=30, bottom=30, left=40, right=40)
        set_cell_margins(cell2, top=30, bottom=30, left=40, right=40)
        if r_idx == 0:
            set_cell_background(cell1, "F1F5F9")
            set_cell_background(cell2, "F1F5F9")
        else:
            if r_idx % 2 == 1:
                set_cell_background(cell1, "FAFAFA")
                set_cell_background(cell2, "FAFAFA")
        for cell, txt in ((cell1, c1_txt), (cell2, c2_txt)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt)
            r.font.size = Pt(6.5)
            if r_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Right Column: Card 3 - Key Technical Implementations
    p_c3_title = col_right.paragraphs[0]
    p_c3_title.paragraph_format.space_before = Pt(0)
    p_c3_title.paragraph_format.space_after = Pt(2)
    r3 = p_c3_title.add_run("3. TECHNICAL INNOVATIONS & FEATURES")
    r3.font.size = Pt(8.5)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    innovations = [
        ("[UX] Tactile Hardware Control: ", "Single short-press Volume Up activates audio listening; second press captures and uploads. Short-press Volume Down stops TTS readout immediately or cancels pending requests. Long press passes through to Android system volume."),
        ("[PRIVACY] On-Premises Vietnamese ASR: ", "Integrates hynt/Zipformer-30M-RNNT-Streaming-6000h via sherpa-onnx. Audio chunks are decoded in ephemeral RAM and discarded upon turn completion, guaranteeing zero voice surveillance."),
        ("[AI AGENT] Verified Remote Guidance: ", "Unlike standard chat models that hallucinate button coordinates, the guidance engine in v0.4.0 enforces strict verification checkpoints, requiring visual proof before instructing the user to press a button."),
        ("[HARDWARE] ToF LiDAR Sensor Extension: ", "Pre-configured integration for STMicroelectronics VL53L9CX Direct Time-of-Flight sensor with a 54x44 matrix (2,376 distance zones) for millimeter-level obstacle distance detection.")
    ]
    for tag_title, feat_desc in innovations:
        p_i = col_right.add_paragraph()
        p_i.paragraph_format.space_before = Pt(0)
        p_i.paragraph_format.space_after = Pt(2.5)
        p_i.paragraph_format.line_spacing = 1.15
        p_i.paragraph_format.left_indent = Inches(0.12)
        r_it = p_i.add_run("■ " + tag_title)
        r_it.font.size = Pt(7.2)
        r_it.font.bold = True
        r_it.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_ix = p_i.add_run(feat_desc)
        r_ix.font.size = Pt(7.0)
        r_ix.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # Right Column: Card 4 - Rigor & Visual Assets
    p_c4_title = col_right.add_paragraph()
    p_c4_title.paragraph_format.space_before = Pt(5)
    p_c4_title.paragraph_format.space_after = Pt(2)
    r4 = p_c4_title.add_run("4. ENGINEERING RIGOR & VISUAL PROOF")
    r4.font.size = Pt(8.5)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    p_rigor = col_right.add_paragraph()
    p_rigor.paragraph_format.space_before = Pt(0)
    p_rigor.paragraph_format.space_after = Pt(3)
    r_rg = p_rigor.add_run(
        "Verification Suite: 322 automated Python pytest cases (WebRTC, tokens, WorkGate) + 329 Jest mobile tests across 17 suites with zero regressions. Strict zero-secret policy in APK."
    )
    r_rg.font.size = Pt(6.8)
    
    # Screenshots Table (3 images side by side)
    img_table = col_right.add_table(rows=1, cols=3)
    img_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    img_table.autofit = False
    img_table.columns[0].width = Inches(1.0)
    img_table.columns[1].width = Inches(1.0)
    img_table.columns[2].width = Inches(1.5)
    
    img_data = [
        (img_table.cell(0, 0), "target/APEC_mat_kinh_cho_nguoi_khiem_thi/docs/screenshots/assistant.png", "Assistant UI", Inches(0.95)),
        (img_table.cell(0, 1), "target/APEC_mat_kinh_cho_nguoi_khiem_thi/docs/screenshots/settings.png", "Settings UI", Inches(0.95)),
        (img_table.cell(0, 2), "target/APEC_mat_kinh_cho_nguoi_khiem_thi/docs/VL53L9CX_ToF.jpg", "VL53L9CX ToF LiDAR", Inches(1.45))
    ]
    for cell, img_path, caption, img_w in img_data:
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=30, bottom=30, left=20, right=20)
        p_img = cell.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(1)
        run_pic = p_img.add_run()
        run_pic.add_picture(img_path, width=img_w)
        
        p_cap = cell.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(1)
        p_cap.paragraph_format.space_after = Pt(0)
        r_cap = p_cap.add_run(caption)
        r_cap.font.size = Pt(5.8)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Security Callout Box
    t_sec = col_right.add_table(rows=1, cols=1)
    t_sec.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sec.autofit = False
    t_sec.columns[0].width = Inches(3.55)
    cell_sec = t_sec.cell(0, 0)
    set_cell_background(cell_sec, "F1F5F9")
    set_cell_margins(cell_sec, top=40, bottom=40, left=50, right=50)
    set_cell_borders(cell_sec, left={"val": "single", "sz": "18", "color": "0284C7"})
    p_sec = cell_sec.paragraphs[0]
    p_sec.paragraph_format.space_before = Pt(0)
    p_sec.paragraph_format.space_after = Pt(0)
    r_sec = p_sec.add_run("Security Guardrails: Bearer tokens stored in Expo SecureStore validated via SHA-256 hashes. Ephemeral memory buffers prevent data persistence. In-memory WorkGate limits single active task per token.")
    r_sec.font.size = Pt(6.5)
    r_sec.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # 4. Footer
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(6)
    p_foot.paragraph_format.space_after = Pt(0)
    r_f1 = p_foot.add_run("APEC Assistive Technologies Project • Open-Source Engineering for the Visually Impaired")
    r_f1.font.size = Pt(6.0)
    r_f1.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    r_f2 = p_foot.add_run("   |   Target Commit: d9ba1e0 • Research & Open Pilot • A4 Poster Document")
    r_f2.font.size = Pt(6.0)
    r_f2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    
    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")

def build_poster_docx_vi(output_path):
    doc = create_document()
    
    # 1. Header Banner Table
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.8)
    header_table.columns[1].width = Inches(2.57)
    
    c0 = header_table.cell(0, 0)
    c1 = header_table.cell(0, 1)
    for c in (c0, c1):
        set_cell_background(c, "091A36")
        set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(2)
    run_badge = p0.add_run("DỰ ÁN NGHIÊN CỨU TRÍ TUỆ NHÂN TẠO & THỊ GIÁC HỖ TRỢ\n")
    run_badge.font.size = Pt(7)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    
    run_title = p0.add_run("APEC MẮT KÍNH CHO NGƯỜI KHIẾM THỊ\n")
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    run_sub = p0.add_run("Thị Giác AI Đa Phương Thức, Nhận Diện Giọng Nói Cục Bộ & Tương Tác Phi Thị Giác")
    run_sub.font.size = Pt(8.5)
    run_sub.font.color.rgb = RGBColor(0xBA, 0xE6, 0xFD)
    
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    
    meta_text = [
        ("Mã nguồn: ", "thanh1912-ut/APEC_mat_kinh..."),
        ("Git SHA mục tiêu: ", "d9ba1e0 (nhánh main)"),
        ("Bản dựng thử nghiệm: ", "Android Pilot v0.4.0 (Code 8)"),
        ("Động cơ giọng nói: ", "Zipformer RNN-T (sherpa-onnx)"),
        ("Thị giác nhân tạo: ", "Google Gemini 3.8 Flash"),
    ]
    for lbl, val in meta_text:
        r_lbl = p1.add_run(lbl)
        r_lbl.font.size = Pt(6.8)
        r_lbl.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        r_val = p1.add_run(val + "\n")
        r_val.font.size = Pt(6.8)
        r_val.font.bold = True
        r_val.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # 2. Key Metrics Ribbon (4 Columns)
    metrics_table = doc.add_table(rows=1, cols=4)
    metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    metrics_table.autofit = False
    
    col_w = Inches(7.37 / 4.0)
    metrics = [
        ("< 180 ms", "Độ Trễ Nhận Diện STT", "0284C7"),
        ("100% Xúc Giác", "Phím Âm Lượng Vật Lý", "10B981"),
        ("Bảo Mật Tuyệt Đối", "Audio Không Lên Mây", "8B5CF6"),
        ("651 Bài Kiểm Thử", "322 Python + 329 Jest", "F59E0B"),
    ]
    for i, (val, lbl, color) in enumerate(metrics):
        cell = metrics_table.cell(0, i)
        cell.width = col_w
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        set_cell_borders(cell, top={"val": "single", "sz": "24", "color": color},
                               bottom={"val": "single", "sz": "4", "color": "E2E8F0"},
                               left={"val": "single", "sz": "4", "color": "E2E8F0"},
                               right={"val": "single", "sz": "4", "color": "E2E8F0"})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        r_v = p.add_run(val + "\n")
        r_v.font.size = Pt(10)
        r_v.font.bold = True
        r_v.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_l = p.add_run(lbl)
        r_l.font.size = Pt(6.5)
        r_l.font.bold = True
        r_l.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3. Two-Column Content Layout
    layout_table = doc.add_table(rows=1, cols=2)
    layout_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    layout_table.autofit = False
    layout_table.columns[0].width = Inches(3.62)
    layout_table.columns[1].width = Inches(3.62)
    
    col_left = layout_table.cell(0, 0)
    col_right = layout_table.cell(0, 1)
    for c in (col_left, col_right):
        set_cell_margins(c, top=0, bottom=0, left=40, right=40)
    
    # Left Column: Card 1 - Mission & Goals
    p_c1_title = col_left.paragraphs[0]
    p_c1_title.paragraph_format.space_before = Pt(0)
    p_c1_title.paragraph_format.space_after = Pt(2)
    r1 = p_c1_title.add_run("1. SỨ MỆNH, MỤC TIÊU & PHẠM VI DỰ ÁN")
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    p_c1_desc = col_left.add_paragraph()
    p_c1_desc.paragraph_format.space_after = Pt(3)
    p_c1_desc.paragraph_format.line_spacing = 1.15
    r_desc = p_c1_desc.add_run(
        "APEC Mắt Kính là giải pháp trợ năng thông minh được thiết kế chuyên biệt cho người khiếm thị, hỗ trợ định vị không gian, nhận diện vật thể và hướng dẫn thao tác thiết bị gia dụng từng bước an toàn."
    )
    r_desc.font.size = Pt(7.4)
    
    bullets_c1 = [
        ("Tương Tác Phi Thị Giác: ", "Thay thế màn hình cảm ứng bằng phím âm lượng vật lý (Tăng = Nghe/Chụp/Gửi; Giảm = Dừng/Hủy) kết hợp âm báo định hướng (earcons)."),
        ("Bảo Vệ Quyền Riêng Tư Cục Bộ: ", "STT tiếng Việt chạy cục bộ bằng mô hình Zipformer RNN-T (~51MB ONNX). Dữ liệu âm thanh thô không bao giờ truyền ra đám mây."),
        ("Hướng Dẫn Thiết Bị Remote (v0.4.0): ", "Quy trình đối thoại thông minh giúp người khiếm thị xác định và bấm đúng nút trên điều khiển máy lạnh, TV và thiết bị gia đình."),
        ("Lộ Trình Cảm Biến Không Gian: ", "Kiến trúc sẵn sàng tích hợp cảm biến Direct ToF LiDAR STMicroelectronics VL53L9CX đo khoảng cách 3D phát hiện vật cản.")
    ]
    for b_title, b_text in bullets_c1:
        p_b = col_left.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(2.5)
        p_b.paragraph_format.line_spacing = 1.15
        p_b.paragraph_format.left_indent = Inches(0.12)
        r_bt = p_b.add_run("■ " + b_title)
        r_bt.font.size = Pt(7.2)
        r_bt.font.bold = True
        r_bt.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_bx = p_b.add_run(b_text)
        r_bx.font.size = Pt(7.0)
        r_bx.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # Left Column: Card 2 - Architecture & Pipeline
    p_c2_title = col_left.add_paragraph()
    p_c2_title.paragraph_format.space_before = Pt(5)
    p_c2_title.paragraph_format.space_after = Pt(2)
    r2 = p_c2_title.add_run("2. KIẾN TRÚC HỆ THỐNG & LUỒNG DỮ LIỆU")
    r2.font.size = Pt(8.5)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    steps = [
        ("A. App Mobile", "React Native/Expo chụp JPEG camera sau (<=1280px) và thu âm M4A/AAC qua phím âm lượng."),
        ("B. Cổng FastAPI", "Xác thực Bearer token qua SHA-256, giới hạn 4MiB, điều phối WorkGate trong RAM."),
        ("C. Chép Lời Cục Bộ", "PyAV chuẩn hóa audio mono 16kHz; sherpa-onnx nhận diện Zipformer tiếng Việt trong RAM."),
        ("D. Thị Giác AI", "Ảnh và câu hỏi gửi đến Gemini 3.8 Flash / 3.1 Flash Lite với cờ store=false bảo mật."),
        ("E. Phản Hồi Trợ Năng", "Kết quả JSON có cấu trúc được đọc qua giọng Android TTS và âm báo, không xung đột TalkBack.")
    ]
    for s_step, s_desc in steps:
        p_s = col_left.add_paragraph()
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = Pt(2)
        p_s.paragraph_format.line_spacing = 1.15
        p_s.paragraph_format.left_indent = Inches(0.1)
        r_st = p_s.add_run(f"[{s_step}]: ")
        r_st.font.size = Pt(7.0)
        r_st.font.bold = True
        r_st.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        r_sx = p_s.add_run(s_desc)
        r_sx.font.size = Pt(6.8)
    
    # Left Column: Spec Table
    t_spec = col_left.add_table(rows=6, cols=2)
    t_spec.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_spec.autofit = False
    t_spec.columns[0].width = Inches(1.3)
    t_spec.columns[1].width = Inches(2.2)
    specs = [
        ("Thành Phần", "Công Nghệ & Thông Số"),
        ("Ứng Dụng Mobile", "React Native, Expo SDK 57, TypeScript"),
        ("Máy Chủ Backend", "Python 3.11, FastAPI, PyAV, WorkGate"),
        ("STT Tiếng Việt", "sherpa-onnx (Zipformer RNN-T 51MB)"),
        ("Thị Giác AI", "Google Gemini 3.8 Flash (store=false)"),
        ("Kho Tri Thức", "PostgreSQL + Redis (Khóa CAS)")
    ]
    for r_idx, (c1_txt, c2_txt) in enumerate(specs):
        cell1 = t_spec.cell(r_idx, 0)
        cell2 = t_spec.cell(r_idx, 1)
        set_cell_margins(cell1, top=30, bottom=30, left=40, right=40)
        set_cell_margins(cell2, top=30, bottom=30, left=40, right=40)
        if r_idx == 0:
            set_cell_background(cell1, "F1F5F9")
            set_cell_background(cell2, "F1F5F9")
        else:
            if r_idx % 2 == 1:
                set_cell_background(cell1, "FAFAFA")
                set_cell_background(cell2, "FAFAFA")
        for cell, txt in ((cell1, c1_txt), (cell2, c2_txt)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt)
            r.font.size = Pt(6.5)
            if r_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Right Column: Card 3 - Key Technical Implementations
    p_c3_title = col_right.paragraphs[0]
    p_c3_title.paragraph_format.space_before = Pt(0)
    p_c3_title.paragraph_format.space_after = Pt(2)
    r3 = p_c3_title.add_run("3. TÍNH NĂNG ĐỘT PHÁ & GIẢI PHÁP KỸ THUẬT")
    r3.font.size = Pt(8.5)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    innovations = [
        ("[TRỢ NĂNG] Điều Khiển Phím Âm Lượng: ", "Nhấn ngắn Tăng âm lượng để mở nghe (âm báo); nhấn lần hai chụp và gửi. Nhấn ngắn Giảm âm lượng dừng đọc ngay hoặc hủy lượt đang xử lý. Giữ lâu phím vẫn chỉnh âm lượng Android bình thường."),
        ("[BẢO MẬT] Nhận Diện Giọng Nói Cục Bộ: ", "Tích hợp hynt/Zipformer-30M-RNNT-Streaming-6000h qua sherpa-onnx. Gói âm thanh giải mã trong RAM tạm thời và xóa sạch sau khi nhận diện, tuyệt đối không gửi giọng nói lên bên thứ ba."),
        ("[AI AGENT] Hướng Dẫn Remote Có Xác Minh: ", "Không suy đoán ảo vị trí nút; phiên bản 0.4.0 bắt buộc xác thực mốc sờ nắn và hình ảnh thiết bị trước khi cho phép người khiếm thị thực hiện thao tác bấm tiếp theo."),
        ("[PHẦN CỨNG] Cảm Biến ToF LiDAR Mở Rộng: ", "Chuẩn bị sẵn kết nối cảm biến STMicroelectronics VL53L9CX Direct ToF với ma trận 54x44 (2.376 vùng đo khoảng cách) quét độ sâu 3D cảnh báo va chạm vật cản.")
    ]
    for tag_title, feat_desc in innovations:
        p_i = col_right.add_paragraph()
        p_i.paragraph_format.space_before = Pt(0)
        p_i.paragraph_format.space_after = Pt(2.5)
        p_i.paragraph_format.line_spacing = 1.15
        p_i.paragraph_format.left_indent = Inches(0.12)
        r_it = p_i.add_run("■ " + tag_title)
        r_it.font.size = Pt(7.2)
        r_it.font.bold = True
        r_it.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_ix = p_i.add_run(feat_desc)
        r_ix.font.size = Pt(7.0)
        r_ix.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # Right Column: Card 4 - Rigor & Visual Assets
    p_c4_title = col_right.add_paragraph()
    p_c4_title.paragraph_format.space_before = Pt(5)
    p_c4_title.paragraph_format.space_after = Pt(2)
    r4 = p_c4_title.add_run("4. CHẤT LƯỢNG KỸ THUẬT & HÌNH ẢNH THỰC TẾ")
    r4.font.size = Pt(8.5)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
    
    p_rigor = col_right.add_paragraph()
    p_rigor.paragraph_format.space_before = Pt(0)
    p_rigor.paragraph_format.space_after = Pt(3)
    r_rg = p_rigor.add_run(
        "Quy Chuẩn Kiểm Thử: Đạt 322 bài kiểm thử tự động pytest Python (WebRTC, tokens, WorkGate) + 329 bài Jest mobile qua 17 bộ test suites đạt 100%. Bảo mật tuyệt đối không lưu API key trong APK."
    )
    r_rg.font.size = Pt(6.8)
    
    # Screenshots Table (3 images side by side)
    img_table = col_right.add_table(rows=1, cols=3)
    img_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    img_table.autofit = False
    img_table.columns[0].width = Inches(1.0)
    img_table.columns[1].width = Inches(1.0)
    img_table.columns[2].width = Inches(1.5)
    
    img_data = [
        (img_table.cell(0, 0), "target/APEC_mat_kinh_cho_nguoi_khiem_thi/docs/screenshots/assistant.png", "Màn hình Trợ lý", Inches(0.95)),
        (img_table.cell(0, 1), "target/APEC_mat_kinh_cho_nguoi_khiem_thi/docs/screenshots/settings.png", "Màn hình Cài đặt", Inches(0.95)),
        (img_table.cell(0, 2), "target/APEC_mat_kinh_cho_nguoi_khiem_thi/docs/VL53L9CX_ToF.jpg", "ToF LiDAR VL53L9CX", Inches(1.45))
    ]
    for cell, img_path, caption, img_w in img_data:
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=30, bottom=30, left=20, right=20)
        p_img = cell.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(1)
        run_pic = p_img.add_run()
        run_pic.add_picture(img_path, width=img_w)
        
        p_cap = cell.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(1)
        p_cap.paragraph_format.space_after = Pt(0)
        r_cap = p_cap.add_run(caption)
        r_cap.font.size = Pt(5.8)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    # Security Callout Box
    t_sec = col_right.add_table(rows=1, cols=1)
    t_sec.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sec.autofit = False
    t_sec.columns[0].width = Inches(3.55)
    cell_sec = t_sec.cell(0, 0)
    set_cell_background(cell_sec, "F1F5F9")
    set_cell_margins(cell_sec, top=40, bottom=40, left=50, right=50)
    set_cell_borders(cell_sec, left={"val": "single", "sz": "18", "color": "0284C7"})
    p_sec = cell_sec.paragraphs[0]
    p_sec.paragraph_format.space_before = Pt(0)
    p_sec.paragraph_format.space_after = Pt(0)
    r_sec = p_sec.add_run("An Toàn Kiến Trúc: Token Bearer lưu trong Expo SecureStore và xác thực qua mã băm SHA-256. Dữ liệu đệm trong RAM được xóa ngay sau xử lý. Cổng WorkGate giới hạn 1 tác vụ/token chống quá tải.")
    r_sec.font.size = Pt(6.5)
    r_sec.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # 4. Footer
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(6)
    p_foot.paragraph_format.space_after = Pt(0)
    r_f1 = p_foot.add_run("Dự Án Công Nghệ Trợ Năng APEC • Kỹ thuật nguồn mở vì cộng đồng người khiếm thị")
    r_f1.font.size = Pt(6.0)
    r_f1.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    r_f2 = p_foot.add_run("   |   Commit mục tiêu: d9ba1e0 • Bản thử nghiệm nghiên cứu • Poster Khổ A4")
    r_f2.font.size = Pt(6.0)
    r_f2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    
    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")

os.makedirs("posters", exist_ok=True)
build_poster_docx_en("posters/APEC_Smart_Glasses_Poster_EN.docx")
build_poster_docx_vi("posters/APEC_Mat_Kinh_Poster_VI.docx")

shutil.copy("posters/APEC_Smart_Glasses_Poster_EN.docx", "APEC_Smart_Glasses_Poster_EN.docx")
shutil.copy("posters/APEC_Mat_Kinh_Poster_VI.docx", "APEC_Mat_Kinh_Poster_VI.docx")
print("Word documents generated and copied to root successfully!")
